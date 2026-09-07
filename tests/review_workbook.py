"""Read-only independent source controls. Run from project root."""
import sys,json,time
from pathlib import Path
from decimal import Decimal
import openpyxl
from backend.engine import ingest,Dataset
source=Path(sys.argv[1])
started=time.perf_counter()
book=openpyxl.load_workbook(source,read_only=True,data_only=True)
totals={k:Decimal(0) for k in ['Sale Value','Net Amount','Quantity','Discount Value','CGST Amount','SGST Amount','IGST Amount(INR)','UTGST Amount']}
rows=0;sheets=[];invoices=set();customers=set();materials=set();brands=set();therapies=set()
for sheet in book:
    it=sheet.iter_rows(values_only=True);headers=next(it,());positions={str(v).strip():i for i,v in enumerate(headers)}
    if not all(k in positions for k in totals):sheets.append({'name':sheet.title,'status':'skipped','rows':sheet.max_row});continue
    count=0
    for row in it:
        if not any(v is not None for v in row):continue
        count+=1
        for k in totals:totals[k]+=Decimal(str(row[positions[k]]).replace(',',''))
        for key,target in [('Invoice Number',invoices),('Payer',customers),('Material',materials),('Mat Brand',brands),('Mat Therapy',therapies)]:target.add(str(row[positions[key]]))
    rows+=count;sheets.append({'name':sheet.title,'rows':count,'status':'transactions'})
book.close()
df,info=ingest(source.read_bytes(),source.name);dataset=Dataset([df]);k=dataset.analytics('overview',{})['kpis']
assert rows==45000,(rows,'fixture count')
assert len(df)==rows
for source_key,metric in [('Sale Value','gross'),('Net Amount','net'),('Quantity','quantity')]:assert Decimal(str(k[metric]))==totals[source_key],(metric,k[metric],totals[source_key])
assert k['invoices']==len(invoices)
assert Decimal(str(k['discount']))==totals['Sale Value']-totals['Net Amount']
assert Decimal(str(k['tax']))==sum(totals[t] for t in ['CGST Amount','SGST Amount','IGST Amount(INR)','UTGST Amount'])
report={'source':str(source),'sheets':sheets,'rows':rows,'independent_totals':{k:str(v) for k,v in totals.items()},'unique':{'invoices':len(invoices),'customers':len(customers),'materials':len(materials),'brands':len(brands),'therapies':len(therapies)},'metadata':dataset.metadata(),'quality':dataset.quality([info]),'api_measure_controls':k,'seconds':round(time.perf_counter()-started,2)}
Path('reports/workbook-review.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['metadata','quality']},indent=2))
print('QUALITY',json.dumps(report['quality']))
dataset.close()
