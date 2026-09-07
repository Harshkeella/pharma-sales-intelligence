import io
from decimal import Decimal
import openpyxl
import pytest
from fastapi.testclient import TestClient
from backend.main import app, SESSIONS, destroy
from backend.engine import ingest,Dataset

def workbook(year=2026,brand='Alpha',optional=True):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Any sheet name'
    headers=['Invoice Date','Invoice Number','Payer','Material','Quantity','Sale Value','Net Amount','Mat Brand','Mat Therapy','Plant','Payer Name']
    extras=['Rate Before Discount','Rate After Discount','Discount Value','CGST Amount','SGST Amount','IGST Amount(INR)','UTGST Amount','CGST Rate','SGST RATE','IGST RATE','UTGST RATE','Expiry Date','Mfg. Date','Sales order date','Delivery Date','P.O date','LR Date']
    ws.append(headers+(extras if optional else []))
    for i in range(2):
        row=[f'0{i+1}-04-{year}',f'I{i}','C'+str(i),'M'+str(i),'2','100.00','80.00',brand if i==0 else 'Beta','Cardio' if i==0 else 'Neuro','P'+str(i),'Customer '+str(i)]
        if optional:row += ['50','40','-20','2','2','0','0','2.5','2.5','0','0',f'01-04-{year+1}',f'01-01-{year}',f'01-04-{year}',f'02-04-{year}',f'01-04-{year}',f'02-04-{year}']
        ws.append(row)
    wb.create_sheet('Summary').append(['Do not import this summary'])
    out=io.BytesIO();wb.save(out);return out.getvalue()

def test_financial_controls_and_filters():
    df,info=ingest(workbook(),'one.xlsx');d=Dataset([df]);k=d.analytics('overview',{})['kpis']
    assert (k['gross'],k['net'],k['discount'],k['quantity'],k['invoices'],k['tax'])==(200,160,40,4,2,8)
    assert (k['discount_pct'],k['avg_invoice'],k['realized_price'])==(20,80,40)
    for key,value in [('brand','Alpha'),('therapy','Cardio'),('month','April'),('year','2026'),('customer','C0'),('plant','P0')]:
        result=d.analytics('overview',{key:[value]})['kpis'];assert result['net']==(160 if key in ('year','month') else 80)
    assert d.analytics('overview',{'brand':['Alpha'],'therapy':['Neuro']})['kpis']['rows']==0
    assert not [i for i in info['issues'] if 'formula mismatch' in i['kind']]
    assert df.iloc[0].shelf_days==365
    assert df.iloc[0].delivery_days==1
    assert len(info['sheets'])==2
    for page in ['trends','products','customers','geography','commercial','supply-chain','pharma-risk','tax']:
        assert d.analytics(page,{})['kpis']['net']==160
    d.close()

def test_missing_optional_tax_is_unknown():
    df,_=ingest(workbook(optional=False),'minimal.xlsx');d=Dataset([df])
    assert d.analytics('overview',{})['kpis']['tax'] is None
    assert d.analytics('pharma-risk',{})['durations']['shelf_days']['average'] is None
    d.close()

def test_dates_and_decimal_relationship_failures():
    wb=openpyxl.load_workbook(io.BytesIO(workbook()));ws=wb.active
    ws.cell(2,6,'100.05');ws.cell(2,13,'-19');ws.cell(2,16,'9');ws.cell(2,23,'01-04-2025')
    out=io.BytesIO();wb.save(out);df,info=ingest(out.getvalue(),'bad.xlsx')
    assert any('formula mismatch' in i['kind'] for i in info['issues'])
    assert any(i['kind']=='Date sequence anomaly' for i in info['issues'])

def test_session_lifecycle_duplicates_and_years():
    with TestClient(app) as c:
        sid=c.post('/api/sessions').json()['id'];base='/api/sessions/'+sid
        original=workbook()
        first=c.post(base+'/files',files={'file':('one.xlsx',original)});assert first.status_code==200,first.text
        fid=first.json()['files'][0]['id']
        assert c.post(base+'/files',files={'file':('copy.xlsx',original)}).status_code==409
        second=c.post(base+'/files',files={'file':('two.xlsx',workbook(2027))});assert second.status_code==200
        assert second.json()['options']['fiscal_year']==['2026-27','2027-28']
        assert c.post(base+'/analytics/overview',json={}).json()['kpis']['net']==320
        comparison=c.post(base+'/comparison',json={'years':['2026-27','2027-28']})
        assert comparison.status_code==200,comparison.text
        assert comparison.json()['days']==2 and comparison.json()['variance']['net']==0
        assert c.delete(base+'/files/'+fid).json()['rows']==2
        assert c.post(base+'/analytics/overview',json={}).json()['kpis']['net']==160
        assert c.post(base+'/analytics/overview',json={'filters':{'invalid':['x']}}).status_code==422
        assert c.post(base+'/export/csv',json={}).status_code==200
        assert c.delete(base).json()['cleared']
        assert c.get(base+'/metadata').status_code==404

def test_schema_rejection_and_invalid_required_row():
    wb=openpyxl.Workbook();wb.active.append(['not a transaction']);out=io.BytesIO();wb.save(out)
    with pytest.raises(ValueError,match='Dataset cannot be analyzed'):ingest(out.getvalue(),'wrong.xlsx')
    wb=openpyxl.load_workbook(io.BytesIO(workbook()));wb.active.cell(2,7,'not money');out=io.BytesIO();wb.save(out)
    df,info=ingest(out.getvalue(),'partially-valid.xlsx');assert len(df)==1 and info['rejected']==1

def test_master_conflicts_across_files():
    a,ia=ingest(workbook(),'a.xlsx');b,ib=ingest(workbook(2027,'Changed'),'b.xlsx');d=Dataset([a,b])
    assert any(i['kind']=='Material Master Conflict' for i in d.quality([ia,ib])['issues'])
    d.close()

@pytest.mark.parametrize('column,kind',[(6,'Gross sales'),(7,'Net sales'),(14,'Discount'),(15,'CGST'),(16,'SGST'),(17,'IGST'),(18,'UTGST')])
def test_each_financial_relationship(column,kind):
    wb=openpyxl.load_workbook(io.BytesIO(workbook()));wb.active.cell(2,column,'1234');out=io.BytesIO();wb.save(out)
    _,info=ingest(out.getvalue(),'mismatch.xlsx')
    assert any(i['kind']==kind+' formula mismatch' for i in info['issues'])
